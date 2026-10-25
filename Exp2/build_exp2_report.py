import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Image, Table, TableStyle, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

EXP2_DIR = r"c:\Users\Tanish\OneDrive\Desktop\DC Home Loan\Exp2"
PDF_PATH = os.path.join(EXP2_DIR, "DC_Lab_Exp_2_Multithreading.pdf")
DOCX_PATH = os.path.join(EXP2_DIR, "DC_Lab_Exp_2_Multithreading.docx")

SERVER_CODE = open(os.path.join(EXP2_DIR, "MultiThreadedLoanServer.java"), "r", encoding="utf-8").read()
CLIENT_CODE = open(os.path.join(EXP2_DIR, "LoanClient.java"), "r", encoding="utf-8").read()

SCREENSHOTS_DIR = os.path.join(EXP2_DIR, "screenshots")

# Canvas page border callback matching lab manual PDF box
def draw_page_border(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(colors.black)
    canvas.setLineWidth(1)
    canvas.rect(30, 30, doc.pagesize[0] - 60, doc.pagesize[1] - 60)
    canvas.restoreState()

# -------------------------------------------------------------
# 1. BUILD DOCX USING PYTHON-DOCX
# -------------------------------------------------------------
def build_docx():
    print("Generating DOCX report...")
    doc = docx.Document()
    
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    # Title
    p_head = doc.add_paragraph()
    p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_h1 = p_head.add_run("Bhartiya Vidya Bhavans Sardar Patel Institute of Technology\n")
    run_h1.bold = True
    run_h1.font.size = Pt(13)
    run_h1.font.name = 'Arial'
    
    run_h2 = p_head.add_run("Munshi Nagar – Andheri West – 400058")
    run_h2.font.size = Pt(11)
    run_h2.font.name = 'Arial'

    p_div = doc.add_paragraph("—" * 65)
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    meta_table = doc.add_table(rows=1, cols=4)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = meta_table.rows[0].cells
    hdr_cells[0].paragraphs[0].add_run("Name: Atharv patil\n          Tanish parakh\n          Gitesh patil").bold = True
    hdr_cells[1].paragraphs[0].add_run("UID: 2024300165\nUID: 2024300157\nUID: 2024300166").bold = True
    hdr_cells[2].paragraphs[0].add_run("DIVISION: C\n\nBATCH: C1").bold = True
    hdr_cells[3].paragraphs[0].add_run("Branch: C.E\n\nDATE:").bold = True
    
    p_div2 = doc.add_paragraph("—" * 65)
    p_div2.alignment = WD_ALIGN_PARAGRAPH.CENTER

    main_box = doc.add_table(rows=1, cols=1)
    main_box.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = main_box.rows[0].cells[0]
    
    p = cell.paragraphs[0]
    r = p.add_run("Experiment No.: 2\n")
    r.bold = True
    r.font.size = Pt(12)
    
    r_aim = p.add_run("Aim: To implement a Java-based Multithreaded Distributed Banking Home Loan Management System demonstrating multithreading and thread synchronization for concurrent loan applications, account EMI transactions, and real-time client notifications from a common server.\n\n")
    r_aim.font.size = Pt(10.5)

    r_th_title = p.add_run("Theory:\n")
    r_th_title.bold = True
    r_th_title.font.size = Pt(11)

    r_th_body = p.add_run(
        "Multithreading in Distributed Systems is an architecture where a server process spawns multiple execution paths (threads) to handle incoming requests from different clients simultaneously. In modern distributed applications like digital banking portals (SBI, HDFC, ICICI), thousands of user transactions occur concurrently. Without multithreading, a single-threaded server would process client requests sequentially—causing severe latency where one client's execution blocks all subsequent clients.\n\n"
        "Why Multithreading? Spawning a dedicated ClientHandler thread for each TCP connection (serverSocket.accept()) isolates client processing, improves system throughput, and ensures high responsiveness.\n\n"
        "Real-Life Examples:\n"
        "• Banking Portals: Multiple account holders pay loan EMIs or submit applications simultaneously without blocking each other.\n"
        "• WhatsApp Web & Slack: Each user sends and receives messages independently via concurrent socket threads.\n"
        "• Zoom & Microsoft Teams: Audio/video streams per participant are processed in separate execution threads.\n\n"
        "Thread Synchronization & Concurrency Control: When multiple threads access shared memory (such as loan registries or account balances), race conditions can corrupt banking state. Using Java synchronized blocks ensures mutual exclusion so that critical operations—such as loan ID generation and EMI balance updates—remain thread-safe.\n\n"
    )
    r_th_body.font.size = Pt(10)

    r_kc = p.add_run("Key Concepts Demonstrated:\n")
    r_kc.bold = True
    r_kc.font.size = Pt(11)
    
    r_kc_body = p.add_run(
        "● Multithreaded Server: Accepts multiple socket connections and handles each in a separate thread.\n"
        "● Thread Synchronization: Uses synchronized locks to enforce concurrency control on shared account balances.\n"
        "● Multithreaded Client: Employs a main thread for user CLI input and a background listener thread for async server responses.\n"
        "● Real-Time Asynchronous Broadcast: Server broadcasts real-time transaction updates to all connected client threads.\n"
        "● Banking Transaction Management: Handles concurrent home loan applications (SUBMIT) and EMI payments (PAY_EMI).\n\n"
    )
    r_kc_body.font.size = Pt(10)

    r_arch = p.add_run("Problem-solving & Architecture Diagram:\n")
    r_arch.bold = True
    r_arch.font.size = Pt(11)
    
    p_img = cell.add_paragraph()
    arch_p = os.path.join(EXP2_DIR, "architecture_diagram.png")
    if os.path.exists(arch_p):
        p_img.add_run().add_picture(arch_p, width=Inches(6.2))
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_code = cell.add_paragraph()
    r_code_t = p_code.add_run("\nCode Implementation:\n")
    r_code_t.bold = True
    r_code_t.font.size = Pt(11)

    r_c1 = p_code.add_run("1. MultiThreadedLoanServer.java\n")
    r_c1.bold = True
    r_c1.font.color.rgb = RGBColor(0, 102, 0)
    
    p_srv = cell.add_paragraph()
    r_srv_text = p_srv.add_run(SERVER_CODE)
    r_srv_text.font.name = 'Consolas'
    r_srv_text.font.size = Pt(8.5)

    p_c2 = cell.add_paragraph()
    r_c2_t = p_c2.add_run("\n2. LoanClient.java\n")
    r_c2_t.bold = True
    r_c2_t.font.color.rgb = RGBColor(0, 102, 0)

    p_cli = cell.add_paragraph()
    r_cli_text = p_cli.add_run(CLIENT_CODE)
    r_cli_text.font.name = 'Consolas'
    r_cli_text.font.size = Pt(8.5)

    p_shots_h = cell.add_paragraph()
    r_sh = p_shots_h.add_run("\nOutput Screenshots:\n")
    r_sh.bold = True
    r_sh.font.size = Pt(11)

    shot_items = [
        ("01_Server_Started.png", "Figure 1: Multithreaded Server Started and Listening on Port 5000"),
        ("04_Concurrent_Loan_Submission.png", "Figure 2: Client 1 Submitting Home Loan Application (HL1001)"),
        ("05_Concurrent_Transaction_Payment.png", "Figure 3: Client 1 Executing Synchronized EMI Payment Transaction"),
        ("06_RealTime_Broadcast_Notification.png", "Figure 4: Client 2 Receiving Real-Time Async Broadcast Notification"),
        ("07_Server_Multithread_Log.png", "Figure 5: Multithreaded Server Execution Log Showing Dedicated Client Threads ([THREAD id])")
    ]

    for fname, caption in shot_items:
        img_path = os.path.join(SCREENSHOTS_DIR, fname)
        if os.path.exists(img_path):
            p_cap = cell.add_paragraph()
            r_cap = p_cap.add_run(caption)
            r_cap.bold = True
            r_cap.font.size = Pt(10)
            
            p_simg = cell.add_paragraph()
            p_simg.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_simg.add_run().add_picture(img_path, width=Inches(6.2))

    p_conc = cell.add_paragraph()
    r_conc_t = p_conc.add_run("\nConclusion:\n")
    r_conc_t.bold = True
    r_conc_t.font.size = Pt(11)
    
    r_conc_b = p_conc.add_run(
        "In this experiment, we successfully implemented a multithreaded distributed banking home loan management system in Java. We demonstrated how spawning dedicated ClientHandler threads for incoming client socket connections enables high parallelism and non-blocking operation. Furthermore, we implemented strict concurrency control using Java synchronized blocks to safeguard shared account databases and loan records against race conditions. The experiment effectively highlighted the critical role of multithreading, thread safety, and real-time asynchronous broadcasting in modern distributed banking software.\n"
    )
    r_conc_b.font.size = Pt(10)

    doc.save(DOCX_PATH)
    print("DOCX build complete:", DOCX_PATH)

# -------------------------------------------------------------
# 2. BUILD PDF USING REPORTLAB
# -------------------------------------------------------------
def build_pdf():
    print("Generating PDF report...")
    doc = SimpleDocTemplate(
        PDF_PATH,
        pagesize=letter,
        rightMargin=45, leftMargin=45, topMargin=45, bottomMargin=45
    )
    
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'HeaderTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        alignment=1,
        textColor=colors.HexColor("#000000")
    )
    
    subtitle_style = ParagraphStyle(
        'HeaderSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=11,
        leading=14,
        alignment=1,
        textColor=colors.HexColor("#333333")
    )
    
    meta_style = ParagraphStyle(
        'MetaText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13.5
    )

    section_heading = ParagraphStyle(
        'SecHeading',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=15,
        textColor=colors.HexColor("#000000"),
        spaceBefore=10,
        spaceAfter=4
    )

    body_style = ParagraphStyle(
        'BodyTextCustom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor("#111111")
    )

    code_style = ParagraphStyle(
        'CodeStyle',
        parent=styles['Normal'],
        fontName='Courier',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor("#002B36")
    )

    story = []

    # College Header
    story.append(Paragraph("Bhartiya Vidya Bhavans Sardar Patel Institute of Technology", title_style))
    story.append(Spacer(1, 2))
    story.append(Paragraph("Munshi Nagar – Andheri West – 400058", subtitle_style))
    story.append(Spacer(1, 6))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.black, spaceBefore=2, spaceAfter=8))

    # Metadata Block Table
    meta_table_data = [
        [
            Paragraph("<b>Name:</b> Atharv patil<br/>&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;<b>Tanish parakh</b><br/>&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;<b>Gitesh patil</b>", meta_style),
            Paragraph("<b>UID:</b> 2024300165<br/><b>UID: 2024300157</b><br/><b>UID:</b> 2024300166", meta_style),
            Paragraph("<b>DIVISION:</b> C<br/><br/><b>BATCH:</b> C1", meta_style),
            Paragraph("<b>Branch:</b> C.E<br/><br/><b>DATE:</b>", meta_style)
        ]
    ]
    t_meta = Table(meta_table_data, colWidths=[160, 130, 110, 122])
    t_meta.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(t_meta)
    story.append(HRFlowable(width="100%", thickness=1, color=colors.black, spaceBefore=6, spaceAfter=10))

    # Exp Title & Aim
    story.append(Paragraph("<b>Experiment No.: 2</b>", section_heading))
    story.append(Spacer(1, 4))
    story.append(Paragraph("<b>Aim:</b> To implement a Java-based Multithreaded Distributed Banking Home Loan Management System demonstrating multithreading and thread synchronization for concurrent loan applications, account EMI transactions, and real-time client notifications from a common server.", body_style))
    story.append(Spacer(1, 8))

    # Theory
    story.append(Paragraph("<b>Theory:</b>", section_heading))
    story.append(Paragraph(
        "<b>Multithreading in Distributed Systems</b> is an architecture where a server process spawns multiple execution paths (threads) to handle incoming requests from different clients simultaneously. In modern distributed applications like digital banking portals (SBI, HDFC, ICICI), thousands of user transactions occur concurrently. Without multithreading, a single-threaded server would process client requests sequentially—causing severe latency where one client's execution blocks all subsequent clients.<br/><br/>"
        "<b>Why Multithreading?</b> Spawning a dedicated <code>ClientHandler</code> thread for each TCP connection (`serverSocket.accept()`) isolates client processing, improves system throughput, and ensures high responsiveness.<br/><br/>"
        "<b>Real-Life Examples:</b><br/>"
        "• <b>Banking Portals:</b> Multiple account holders pay loan EMIs or submit applications simultaneously without blocking each other.<br/>"
        "• <b>WhatsApp Web & Slack:</b> Each user sends and receives messages independently via concurrent socket threads.<br/>"
        "• <b>Zoom & Microsoft Teams:</b> Audio/video streams per participant are processed in separate execution threads.<br/><br/>"
        "<b>Thread Synchronization & Concurrency Control:</b> When multiple threads access shared memory (such as loan registries or account balances), race conditions can corrupt banking state. Using Java <code>synchronized</code> blocks ensures mutual exclusion so that critical operations—such as loan ID generation and EMI balance updates—remain thread-safe.",
        body_style
    ))
    story.append(Spacer(1, 8))

    # Key Concepts
    story.append(Paragraph("<b>Key Concepts Demonstrated:</b>", section_heading))
    story.append(Paragraph(
        "• <b>Multithreaded Server:</b> Accepts multiple socket connections and handles each in a separate thread.<br/>"
        "• <b>Thread Synchronization:</b> Uses <code>synchronized</code> locks to enforce concurrency control on shared account balances.<br/>"
        "• <b>Multithreaded Client:</b> Employs a main thread for user CLI input and a background listener thread for async server responses.<br/>"
        "• <b>Real-Time Asynchronous Broadcast:</b> Server broadcasts real-time transaction updates to all connected client threads.<br/>"
        "• <b>Banking Transaction Management:</b> Handles concurrent home loan applications (`SUBMIT`) and EMI payments (`PAY_EMI`).",
        body_style
    ))
    story.append(Spacer(1, 10))

    # Architecture Diagram
    story.append(Paragraph("<b>Problem-solving & System Architecture:</b>", section_heading))
    story.append(Spacer(1, 4))
    arch_img_path = os.path.join(EXP2_DIR, "architecture_diagram.png")
    if os.path.exists(arch_img_path):
        story.append(Image(arch_img_path, width=480, height=275))
    story.append(Spacer(1, 10))

    # Code Section
    story.append(Paragraph("<b>Code Implementation:</b>", section_heading))
    story.append(Spacer(1, 4))
    story.append(Paragraph("<b>1. MultiThreadedLoanServer.java (Server with Concurrency Control)</b>", ParagraphStyle('SubHead', parent=body_style, fontName='Helvetica-Bold', textColor=colors.HexColor("#006600"))))
    story.append(Spacer(1, 4))

    def create_code_paragraphs(code_text):
        lines = code_text.splitlines()
        flowables = []
        chunk = []
        for line in lines:
            chunk.append(line)
            if len(chunk) >= 40:
                chunk_str = "\n".join(chunk).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>").replace(" ", "&nbsp;")
                flowables.append(Paragraph(chunk_str, code_style))
                flowables.append(Spacer(1, 4))
                chunk = []
        if chunk:
            chunk_str = "\n".join(chunk).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>").replace(" ", "&nbsp;")
            flowables.append(Paragraph(chunk_str, code_style))
        return flowables

    story.extend(create_code_paragraphs(SERVER_CODE))
    story.append(Spacer(1, 10))

    story.append(Paragraph("<b>2. LoanClient.java (Multithreaded Client with Listener Thread)</b>", ParagraphStyle('SubHead2', parent=body_style, fontName='Helvetica-Bold', textColor=colors.HexColor("#006600"))))
    story.append(Spacer(1, 4))
    story.extend(create_code_paragraphs(CLIENT_CODE))
    story.append(Spacer(1, 12))

    # Output Screenshots Section
    story.append(Paragraph("<b>Output Screenshots:</b>", section_heading))
    story.append(Spacer(1, 6))

    shot_items = [
        ("01_Server_Started.png", "Figure 1: Multithreaded Server Started and Listening on Port 5000"),
        ("04_Concurrent_Loan_Submission.png", "Figure 2: Client 1 Submitting Home Loan Application (HL1001)"),
        ("05_Concurrent_Transaction_Payment.png", "Figure 3: Client 1 Executing Synchronized EMI Payment Transaction (Balance updated)"),
        ("06_RealTime_Broadcast_Notification.png", "Figure 4: Client 2 Receiving Real-Time Async Broadcast Notification"),
        ("07_Server_Multithread_Log.png", "Figure 5: Multithreaded Server Execution Log Showing Dedicated Client Threads ([THREAD id])")
    ]

    for fname, caption in shot_items:
        img_p = os.path.join(SCREENSHOTS_DIR, fname)
        if os.path.exists(img_p):
            story.append(Paragraph(f"<b>{caption}</b>", body_style))
            story.append(Spacer(1, 4))
            story.append(Image(img_p, width=480, height=270))
            story.append(Spacer(1, 10))

    # Conclusion
    story.append(Paragraph("<b>Conclusion:</b>", section_heading))
    story.append(Paragraph(
        "In this experiment, we successfully implemented a multithreaded distributed banking home loan management system in Java. We demonstrated how spawning dedicated <code>ClientHandler</code> threads for incoming client socket connections enables high parallelism and non-blocking operation. Furthermore, we implemented strict concurrency control using Java <code>synchronized</code> blocks to safeguard shared account databases and loan records against race conditions. The experiment effectively highlighted the critical role of multithreading, thread safety, and real-time asynchronous broadcasting in modern distributed banking software.",
        body_style
    ))

    doc.build(story, onFirstPage=draw_page_border, onLaterPages=draw_page_border)
    print("PDF build complete:", PDF_PATH)

if __name__ == "__main__":
    build_docx()
    build_pdf()
