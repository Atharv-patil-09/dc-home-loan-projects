import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Image, Table, TableStyle, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

EXP2_DIR = r"c:\Users\Tanish\OneDrive\Desktop\DC Home Loan\Exp2"
PDF_PATH = os.path.join(EXP2_DIR, "DC_Lab_Exp_2_Multithreading.pdf")
DOCX_PATH = os.path.join(EXP2_DIR, "DC_Lab_Exp_2_Multithreading.docx")

SERVER_CODE = open(os.path.join(EXP2_DIR, "MultiThreadedLoanServer.java"), "r", encoding="utf-8").read()
CLIENT_CODE = open(os.path.join(EXP2_DIR, "LoanClient.java"), "r", encoding="utf-8").read()

SCREENSHOTS_DIR = os.path.join(EXP2_DIR, "screenshots")
LOGO_PATH = os.path.join(EXP2_DIR, "spit_logo.png")
ARCH_PATH = os.path.join(EXP2_DIR, "architecture_diagram.png")

# Canvas page border callback matching lab manual PDF box
def draw_page_border(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(colors.black)
    canvas.setLineWidth(1)
    canvas.rect(30, 30, doc.pagesize[0] - 60, doc.pagesize[1] - 60)
    canvas.restoreState()

# -------------------------------------------------------------
# 1. BUILD EXACT SPIT TEMPLATE DOCX
# -------------------------------------------------------------
def build_exact_spit_docx():
    print("Generating Exact SPIT Template DOCX report...")
    doc = docx.Document()
    
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    # Header Table: Logo on Left, Title on Right
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_cells = header_table.rows[0].cells
    
    # Left Cell: Logo
    p_logo = header_cells[0].paragraphs[0]
    if os.path.exists(LOGO_PATH):
        p_logo.add_run().add_picture(LOGO_PATH, width=Inches(1.2))
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Right Cell: College Name
    p_head = header_cells[1].paragraphs[0]
    p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p_head.add_run("Bhartiya Vidya Bhavans\n")
    r1.bold = True
    r1.font.size = Pt(13)
    r1.font.name = 'Arial'
    
    r2 = p_head.add_run("Sardar Patel Institute of Technology\n")
    r2.bold = True
    r2.font.size = Pt(13)
    r2.font.name = 'Arial'

    r3 = p_head.add_run("Munshi Nagar – Andheri West – 400058")
    r3.font.size = Pt(11)
    r3.font.name = 'Arial'

    # Divider Line
    p_div = doc.add_paragraph("—" * 65)
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Student Metadata Table
    meta_table = doc.add_table(rows=1, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = meta_table.rows[0].cells
    
    hdr_cells[0].paragraphs[0].add_run("Name: Atharv patil (UID: 2024300165)\n          Tanish parakh (UID: 2024300157)\n          Gitesh patil (UID: 2024300166)\nBranch: C.E").bold = True
    hdr_cells[1].paragraphs[0].add_run("DIVISION: C\nBATCH: C1\nDATE: 11/08/2026").bold = True
    
    p_div2 = doc.add_paragraph("—" * 65)
    p_div2.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Main Template Table with specific section rows (matching PDF template layout)
    table = doc.add_table(rows=9, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Row 0: Experiment No.
    c0 = table.rows[0].cells[0]
    p0 = c0.paragraphs[0]
    r0 = p0.add_run("Experiment No.: 2")
    r0.bold = True
    r0.font.size = Pt(11)

    # Row 1: Aim
    c1 = table.rows[1].cells[0]
    p1 = c1.paragraphs[0]
    r1_title = p1.add_run("Aim : ")
    r1_title.bold = True
    r1_title.font.size = Pt(11)
    r1_body = p1.add_run("To implement a Java-based Multithreaded Distributed Banking Home Loan Management System demonstrating multithreading and thread synchronization for concurrent loan applications, account EMI transactions, and real-time client notifications from a common server.")
    r1_body.font.size = Pt(10)

    # Row 2: Theory
    c2 = table.rows[2].cells[0]
    p2 = c2.paragraphs[0]
    r2_title = p2.add_run("Theory :\n")
    r2_title.bold = True
    r2_title.font.size = Pt(11)
    r2_body = p2.add_run(
        "Multithreading in Distributed Systems is a design pattern where a server process spawns multiple concurrent execution paths (threads) to process incoming requests from different clients simultaneously. In modern distributed banking portals (SBI, HDFC, ICICI), thousands of user transactions occur concurrently. Without multithreading, a single-threaded server processes client requests sequentially—causing severe latency where one client's execution blocks all subsequent clients.\n\n"
        "Why Multithreading?\n"
        "• Spawning a dedicated ClientHandler thread for each TCP connection (serverSocket.accept()) isolates client processing, improves system throughput, and ensures high responsiveness.\n\n"
        "Real-Life Examples:\n"
        "• Banking Portals: Multiple account holders pay loan EMIs or submit applications simultaneously without blocking each other.\n"
        "• WhatsApp Web & Slack: Each user sends and receives messages independently via concurrent socket threads.\n"
        "• Zoom & Microsoft Teams: Audio/video streams per participant are processed in separate execution threads.\n\n"
        "Thread Synchronization & Concurrency Control:\n"
        "When multiple threads access shared memory (such as loan registries or account balances), race conditions can corrupt banking state. Using Java synchronized blocks ensures mutual exclusion so that critical operations—such as loan ID generation and EMI balance updates—remain thread-safe.\n\n"
        "Key Concepts Demonstrated:\n"
        "● Client: Connects to multithreaded server and sends loan requests.\n"
        "● Server: Spawns a dedicated thread per client and manages synchronized banking state.\n"
        "● Multithreading: Enables concurrent, non-blocking execution paths.\n"
        "● Concurrency Control: Protects shared account balances against race conditions.\n"
        "● Real-time Broadcast: Notifies all connected client threads asynchronously of banking transactions."
    )
    r2_body.font.size = Pt(10)

    # Row 3: Problem-solving
    c3 = table.rows[3].cells[0]
    p3 = c3.paragraphs[0]
    r3_title = p3.add_run("Problem-solving :\n")
    r3_title.bold = True
    r3_title.font.size = Pt(11)
    
    if os.path.exists(ARCH_PATH):
        p3_img = c3.add_paragraph()
        p3_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p3_img.add_run().add_picture(ARCH_PATH, width=Inches(6.0))

    # Row 4: Code
    c4 = table.rows[4].cells[0]
    p4 = c4.paragraphs[0]
    r4_title = p4.add_run("Code :\n")
    r4_title.bold = True
    r4_title.font.size = Pt(11)

    r4_c1 = p4.add_run("1. MultiThreadedLoanServer.java\n")
    r4_c1.bold = True
    r4_c1.font.color.rgb = RGBColor(0, 102, 0)

    p4_srv = c4.add_paragraph()
    r4_srv_text = p4_srv.add_run(SERVER_CODE)
    r4_srv_text.font.name = 'Consolas'
    r4_srv_text.font.size = Pt(8.5)

    p4_c2 = c4.add_paragraph()
    r4_c2_t = p4_c2.add_run("\n2. LoanClient.java\n")
    r4_c2_t.bold = True
    r4_c2_t.font.color.rgb = RGBColor(0, 102, 0)

    p4_cli = c4.add_paragraph()
    r4_cli_text = p4_cli.add_run(CLIENT_CODE)
    r4_cli_text.font.name = 'Consolas'
    r4_cli_text.font.size = Pt(8.5)

    # Row 5: Output screenshot
    c5 = table.rows[5].cells[0]
    p5 = c5.paragraphs[0]
    r5_title = p5.add_run("Output screenshot :\n")
    r5_title.bold = True
    r5_title.font.size = Pt(11)

    shot_items = [
        ("01_Server_Started.png", "Figure 1: Multithreaded Server Started and Listening on Port 5000"),
        ("04_Concurrent_Loan_Submission.png", "Figure 2: Client 1 Submitting Home Loan Application (HL1001)"),
        ("05_Concurrent_Transaction_Payment.png", "Figure 3: Client 1 Executing Synchronized EMI Payment Transaction"),
        ("06_RealTime_Broadcast_Notification.png", "Figure 4: Client 2 Receiving Real-Time Async Broadcast Notification"),
        ("07_Server_Multithread_Log.png", "Figure 5: Multithreaded Server Execution Log Showing Dedicated Client Threads ([THREAD id])")
    ]

    for fname, caption in shot_items:
        img_path = os.path.join(SCREENSHOTS_DIR, fname)
        if os.path.exists(img_path):
            p_cap = c5.add_paragraph()
            r_cap = p_cap.add_run(caption)
            r_cap.bold = True
            r_cap.font.size = Pt(9.5)
            
            p_simg = c5.add_paragraph()
            p_simg.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_simg.add_run().add_picture(img_path, width=Inches(6.0))

    # Row 6: Conclusion
    c6 = table.rows[6].cells[0]
    p6 = c6.paragraphs[0]
    r6_title = p6.add_run("Conclusion:\n")
    r6_title.bold = True
    r6_title.font.size = Pt(11)
    
    r6_body = p6.add_run(
        "In this experiment, we successfully implemented a multithreaded distributed banking home loan management system in Java. We demonstrated how spawning dedicated ClientHandler threads for incoming client socket connections enables high parallelism and non-blocking operation. Furthermore, we implemented strict concurrency control using Java synchronized blocks to safeguard shared account databases and loan records against race conditions. The experiment effectively highlighted the critical role of multithreading, thread safety, and real-time asynchronous broadcasting in modern distributed banking software."
    )
    r6_body.font.size = Pt(10)

    # Row 7: Research Paper note
    c7 = table.rows[7].cells[0]
    p7 = c7.paragraphs[0]
    r7_text = p7.add_run("Attach research paper based on experiment performed into the classroom (Not for all the practicals)")
    r7_text.font.size = Pt(9.5)
    r7_text.italic = True

    # Row 8: Document Submission note
    c8 = table.rows[8].cells[0]
    p8 = c8.paragraphs[0]
    r8_text = p8.add_run("(Students are required to prepare one doc file for each experiment and then upload it in pdf form.)")
    r8_text.font.size = Pt(9.5)
    r8_text.italic = True

    doc.save(DOCX_PATH)
    print("Exact SPIT Template DOCX build complete:", DOCX_PATH)

# -------------------------------------------------------------
# 2. BUILD EXACT SPIT TEMPLATE PDF
# -------------------------------------------------------------
def build_exact_spit_pdf():
    print("Generating Exact SPIT Template PDF report...")
    doc = SimpleDocTemplate(
        PDF_PATH,
        pagesize=letter,
        rightMargin=45, leftMargin=45, topMargin=45, bottomMargin=45
    )
    
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'HeaderTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        alignment=1,
        textColor=colors.HexColor("#000000")
    )
    
    subtitle_style = ParagraphStyle(
        'HeaderSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=11,
        leading=14,
        alignment=1,
        textColor=colors.HexColor("#333333")
    )
    
    meta_style = ParagraphStyle(
        'MetaText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13.5
    )

    section_heading = ParagraphStyle(
        'SecHeading',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=15,
        textColor=colors.HexColor("#000000"),
        spaceBefore=6,
        spaceAfter=4
    )

    body_style = ParagraphStyle(
        'BodyTextCustom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor("#111111")
    )

    code_style = ParagraphStyle(
        'CodeStyle',
        parent=styles['Normal'],
        fontName='Courier',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor("#002B36")
    )

    story = []

    # Header Table: Logo + Text
    logo_img = Image(LOGO_PATH, width=70, height=70) if os.path.exists(LOGO_PATH) else Paragraph("", body_style)
    text_p = Paragraph("<b>Bhartiya Vidya Bhavans</b><br/><b>Sardar Patel Institute of Technology</b><br/>Munshi Nagar – Andheri West – 400058", title_style)
    
    head_table = Table([[logo_img, text_p]], colWidths=[90, 430])
    head_table.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('ALIGN', (0,0), (0,0), 'CENTER'),
        ('ALIGN', (1,0), (1,0), 'CENTER'),
    ]))
    story.append(head_table)
    story.append(Spacer(1, 4))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.black, spaceBefore=2, spaceAfter=6))

    # Metadata Block Table
    meta_table_data = [
        [
            Paragraph("<b>Name:</b> Atharv patil (2024300165)<br/>&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;<b>Tanish parakh (2024300157)</b><br/>&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;<b>Gitesh patil (2024300166)</b><br/><b>Branch:</b> C.E", meta_style),
            Paragraph("<b>DIVISION:</b> C<br/><b>BATCH:</b> C1<br/><b>DATE:</b> 11/08/2026", meta_style)
        ]
    ]
    t_meta = Table(meta_table_data, colWidths=[340, 180])
    t_meta.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(t_meta)
    story.append(HRFlowable(width="100%", thickness=1, color=colors.black, spaceBefore=4, spaceAfter=8))

    # Section 1: Experiment No.
    story.append(Paragraph("<b>Experiment No.: 2</b>", section_heading))
    story.append(Spacer(1, 4))

    # Section 2: Aim
    story.append(Paragraph("<b>Aim :</b> To implement a Java-based Multithreaded Distributed Banking Home Loan Management System demonstrating multithreading and thread synchronization for concurrent loan applications, account EMI transactions, and real-time client notifications from a common server.", body_style))
    story.append(Spacer(1, 8))

    # Section 3: Theory
    story.append(Paragraph("<b>Theory :</b>", section_heading))
    story.append(Paragraph(
        "<b>Multithreading in Distributed Systems</b> is an architecture where a server process spawns multiple execution paths (threads) to handle incoming requests from different clients simultaneously. In modern distributed applications like digital banking portals (SBI, HDFC, ICICI), thousands of user transactions occur concurrently. Without multithreading, a single-threaded server would process client requests sequentially—causing severe latency where one client's execution blocks all subsequent clients.<br/><br/>"
        "<b>Why Multithreading?</b> Spawning a dedicated <code>ClientHandler</code> thread for each TCP connection (`serverSocket.accept()`) isolates client processing, improves system throughput, and ensures high responsiveness.<br/><br/>"
        "<b>Real-Life Examples:</b><br/>"
        "• <b>Banking Portals:</b> Multiple account holders pay loan EMIs or submit applications simultaneously without blocking each other.<br/>"
        "• <b>WhatsApp Web & Slack:</b> Each user sends and receives messages independently via concurrent socket threads.<br/>"
        "• <b>Zoom & Microsoft Teams:</b> Audio/video streams per participant are processed in separate execution threads.<br/><br/>"
        "<b>Thread Synchronization & Concurrency Control:</b> When multiple threads access shared memory (such as loan registries or account balances), race conditions can corrupt banking state. Using Java <code>synchronized</code> blocks ensures mutual exclusion so that critical operations—such as loan ID generation and EMI balance updates—remain thread-safe.<br/><br/>"
        "<b>Key Concepts Demonstrated:</b><br/>"
        "● <b>Client:</b> Connects to multithreaded server and sends loan requests.<br/>"
        "● <b>Server:</b> Spawns a dedicated thread per client and manages synchronized banking state.<br/>"
        "● <b>Multithreading:</b> Enables concurrent, non-blocking execution paths.<br/>"
        "● <b>Concurrency Control:</b> Protects shared account balances against race conditions.<br/>"
        "● <b>Real-time Broadcast:</b> Notifies all connected client threads asynchronously of banking transactions.",
        body_style
    ))
    story.append(Spacer(1, 10))

    # Section 4: Problem-solving & Architecture
    story.append(Paragraph("<b>Problem-solving :</b>", section_heading))
    story.append(Spacer(1, 4))
    if os.path.exists(ARCH_PATH):
        story.append(Image(ARCH_PATH, width=480, height=275))
    story.append(Spacer(1, 10))

    # Section 5: Code
    story.append(Paragraph("<b>Code :</b>", section_heading))
    story.append(Spacer(1, 4))
    story.append(Paragraph("<b>1. MultiThreadedLoanServer.java (Server with Concurrency Control)</b>", ParagraphStyle('SubHead', parent=body_style, fontName='Helvetica-Bold', textColor=colors.HexColor("#006600"))))
    story.append(Spacer(1, 4))

    def create_code_paragraphs(code_text):
        lines = code_text.splitlines()
        flowables = []
        chunk = []
        for line in lines:
            chunk.append(line)
            if len(chunk) >= 40:
                chunk_str = "\n".join(chunk).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>").replace(" ", "&nbsp;")
                flowables.append(Paragraph(chunk_str, code_style))
                flowables.append(Spacer(1, 4))
                chunk = []
        if chunk:
            chunk_str = "\n".join(chunk).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>").replace(" ", "&nbsp;")
            flowables.append(Paragraph(chunk_str, code_style))
        return flowables

    story.extend(create_code_paragraphs(SERVER_CODE))
    story.append(Spacer(1, 10))

    story.append(Paragraph("<b>2. LoanClient.java (Multithreaded Client with Listener Thread)</b>", ParagraphStyle('SubHead2', parent=body_style, fontName='Helvetica-Bold', textColor=colors.HexColor("#006600"))))
    story.append(Spacer(1, 4))
    story.extend(create_code_paragraphs(CLIENT_CODE))
    story.append(Spacer(1, 12))

    # Section 6: Output screenshot
    story.append(Paragraph("<b>Output screenshot :</b>", section_heading))
    story.append(Spacer(1, 6))

    shot_items = [
        ("01_Server_Started.png", "Figure 1: Multithreaded Server Started and Listening on Port 5000"),
        ("04_Concurrent_Loan_Submission.png", "Figure 2: Client 1 Submitting Home Loan Application (HL1001)"),
        ("05_Concurrent_Transaction_Payment.png", "Figure 3: Client 1 Executing Synchronized EMI Payment Transaction (Balance updated)"),
        ("06_RealTime_Broadcast_Notification.png", "Figure 4: Client 2 Receiving Real-Time Async Broadcast Notification"),
        ("07_Server_Multithread_Log.png", "Figure 5: Multithreaded Server Execution Log Showing Dedicated Client Threads ([THREAD id])")
    ]

    for fname, caption in shot_items:
        img_p = os.path.join(SCREENSHOTS_DIR, fname)
        if os.path.exists(img_p):
            story.append(Paragraph(f"<b>{caption}</b>", body_style))
            story.append(Spacer(1, 4))
            story.append(Image(img_p, width=480, height=270))
            story.append(Spacer(1, 10))

    # Section 7: Conclusion
    story.append(Paragraph("<b>Conclusion:</b> (It should include what the student has learned/ understood from the experiment).", section_heading))
    story.append(Paragraph(
        "In this experiment, we successfully implemented a multithreaded distributed banking home loan management system in Java. We demonstrated how spawning dedicated <code>ClientHandler</code> threads for incoming client socket connections enables high parallelism and non-blocking operation. Furthermore, we implemented strict concurrency control using Java <code>synchronized</code> blocks to safeguard shared account databases and loan records against race conditions. The experiment effectively highlighted the critical role of multithreading, thread safety, and real-time asynchronous broadcasting in modern distributed banking software.",
        body_style
    ))
    story.append(Spacer(1, 10))

    # Section 8 & 9 Notes
    story.append(Paragraph("<i>Attach research paper based on experiment performed into the classroom (Not for all the practicals)</i>", body_style))
    story.append(Spacer(1, 4))
    story.append(Paragraph("<i>(Students are required to prepare one doc file for each experiment and then upload it in pdf form.)</i>", body_style))

    doc.build(story, onFirstPage=draw_page_border, onLaterPages=draw_page_border)
    print("Exact SPIT Template PDF build complete:", PDF_PATH)

if __name__ == "__main__":
    build_exact_spit_docx()
    build_exact_spit_pdf()
