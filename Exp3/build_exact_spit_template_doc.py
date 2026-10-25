import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Image, Table, TableStyle, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

EXP3_DIR = r"c:\Users\Tanish\OneDrive\Desktop\DC Home Loan\Exp3"
PDF_PATH = os.path.join(EXP3_DIR, "DC_Lab_Exp_3_ClockSync.pdf")
DOCX_PATH = os.path.join(EXP3_DIR, "DC_Lab_Exp_3_ClockSync.docx")

SERVER_CODE = open(os.path.join(EXP3_DIR, "MultiThreadedLoanServer.java"), "r", encoding="utf-8").read()
CLIENT_CODE = open(os.path.join(EXP3_DIR, "LoanClient.java"), "r", encoding="utf-8").read()
LAMPORT_CODE = open(os.path.join(EXP3_DIR, "LamportClockBanking.java"), "r", encoding="utf-8").read()

SCREENSHOTS_DIR = os.path.join(EXP3_DIR, "screenshots")
LOGO_PATH = r"c:\Users\Tanish\OneDrive\Desktop\DC Home Loan\Exp2\spit_logo.png"
ARCH_PATH = os.path.join(EXP3_DIR, "architecture_diagram.png")

def draw_page_border(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(colors.black)
    canvas.setLineWidth(1)
    canvas.rect(30, 30, doc.pagesize[0] - 60, doc.pagesize[1] - 60)
    canvas.restoreState()

# -------------------------------------------------------------
# 1. BUILD EXACT SPIT TEMPLATE DOCX
# -------------------------------------------------------------
def build_exact_spit_docx():
    print("Generating Exact SPIT Template DOCX report for Exp 3...")
    doc = docx.Document()
    
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    # Header Table: Logo + Text
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_cells = header_table.rows[0].cells
    
    p_logo = header_cells[0].paragraphs[0]
    if os.path.exists(LOGO_PATH):
        p_logo.add_run().add_picture(LOGO_PATH, width=Inches(1.2))
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    p_head = header_cells[1].paragraphs[0]
    p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p_head.add_run("Bhartiya Vidya Bhavans\n")
    r1.bold = True
    r1.font.size = Pt(13)
    r1.font.name = 'Arial'
    
    r2 = p_head.add_run("Sardar Patel Institute of Technology\n")
    r2.bold = True
    r2.font.size = Pt(13)
    r2.font.name = 'Arial'

    r3 = p_head.add_run("Munshi Nagar – Andheri West – 400058")
    r3.font.size = Pt(11)
    r3.font.name = 'Arial'

    p_div = doc.add_paragraph("—" * 65)
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Metadata Table
    meta_table = doc.add_table(rows=1, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = meta_table.rows[0].cells
    
    hdr_cells[0].paragraphs[0].add_run("Name: Atharv patil (UID: 2024300165)\n          Tanish parakh (UID: 2024300157)\n          Gitesh patil (UID: 2024300166)\nBranch: C.E").bold = True
    hdr_cells[1].paragraphs[0].add_run("DIVISION: C\nBATCH: C1\nDATE: 23/08/2026").bold = True
    
    p_div2 = doc.add_paragraph("—" * 65)
    p_div2.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Main Template Table with specific section rows
    table = doc.add_table(rows=9, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Row 0: Experiment No.
    c0 = table.rows[0].cells[0]
    p0 = c0.paragraphs[0]
    r0 = p0.add_run("Experiment No.: 3")
    r0.bold = True
    r0.font.size = Pt(11)

    # Row 1: Aim
    c1 = table.rows[1].cells[0]
    p1 = c1.paragraphs[0]
    r1_title = p1.add_run("Aim : ")
    r1_title.bold = True
    r1_title.font.size = Pt(11)
    r1_body = p1.add_run("To extend a Java-based Multithreaded Distributed Banking Home Loan Management System by implementing Cristian's Algorithm (Client-Pull Method) and the Berkeley Algorithm (Server-Push Time Daemon), ensuring precise clock synchronization and accurate timestamping for concurrent loan applications and EMI transactions across distributed clients.")
    r1_body.font.size = Pt(10)

    # Row 2: Theory
    c2 = table.rows[2].cells[0]
    p2 = c2.paragraphs[0]
    r2_title = p2.add_run("Theory :\n")
    r2_title.bold = True
    r2_title.font.size = Pt(11)
    r2_body = p2.add_run(
        "Clock Synchronization in Distributed Banking Systems is critical to align local system times across independent client nodes and servers, maintaining a globally consistent transaction ledger. In distributed banking, if physical clocks drift due to hardware oscillator variance, financial transactions (loan submissions, EMI account debits) can be logged out of order.\n\n"
        "1. Cristian's Algorithm (Client-Pull Method):\n"
        "• Concept: A Loan Client actively requests the exact time from the central Loan Server right before executing an EMI transaction.\n"
        "• Implementation: The client records local time T0 before dispatching the SYNC_CRISTIAN request to the server.\n"
        "• Server Response: The server receives this request and replies with its current system time T_server.\n"
        "• Calculation: The client receives the response at T1, computes Round Trip Time RTT = T1 - T0, and adjusts its clock to T_sync = T_server + (RTT / 2).\n\n"
        "2. Berkeley Algorithm (Server-Push Time Daemon Method):\n"
        "• Concept: The central Loan Server acts as an active Time Daemon, ensuring all connected banking clients are synchronized.\n"
        "• Implementation: A background daemon thread in the server periodically polls active client sockets with a POLL_TIME prompt.\n"
        "• Calculation: Clients respond with their local timestamps (TIME_REPLY). The server calculates the average time skew and sends individual offset adjustments back so clients correct local drift.\n\n"
        "Key Concepts Demonstrated:\n"
        "● Cristian's Algorithm: Adjusts client clock by querying time server and compensating for RTT / 2.\n"
        "● Berkeley Algorithm: Utilizes a central time daemon to poll clients, calculate time skew, and issue offset adjustments.\n"
        "● Clock Synchronization: Aligns local system times across distributed banking nodes.\n"
        "● Transaction Integrity: Guarantees concurrent home loan submissions and EMI payments are recorded with synchronized timestamps."
    )
    r2_body.font.size = Pt(10)

    # Row 3: Problem-solving
    c3 = table.rows[3].cells[0]
    p3 = c3.paragraphs[0]
    r3_title = p3.add_run("Problem-solving :\n")
    r3_title.bold = True
    r3_title.font.size = Pt(11)
    
    if os.path.exists(ARCH_PATH):
        p3_img = c3.add_paragraph()
        p3_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p3_img.add_run().add_picture(ARCH_PATH, width=Inches(6.2))

    # Row 4: Code
    c4 = table.rows[4].cells[0]
    p4 = c4.paragraphs[0]
    r4_title = p4.add_run("Code :\n")
    r4_title.bold = True
    r4_title.font.size = Pt(11)

    r4_c1 = p4.add_run("1. MultiThreadedLoanServer.java (Server with Cristian & Berkeley Daemon)\n")
    r4_c1.bold = True
    r4_c1.font.color.rgb = RGBColor(0, 102, 0)

    p4_code1 = c4.add_paragraph()
    r4_code1_text = p4_code1.add_run(SERVER_CODE)
    r4_code1_text.font.name = 'Consolas'
    r4_code1_text.font.size = Pt(8.5)

    p4_c2 = c4.add_paragraph()
    r4_c2_t = p4_c2.add_run("\n2. LoanClient.java (Client with Cristian Sync & Berkeley Listener)\n")
    r4_c2_t.bold = True
    r4_c2_t.font.color.rgb = RGBColor(0, 102, 0)

    p4_code2 = c4.add_paragraph()
    r4_code2_text = p4_code2.add_run(CLIENT_CODE)
    r4_code2_text.font.name = 'Consolas'
    r4_code2_text.font.size = Pt(8.5)

    # Row 5: Output screenshot
    c5 = table.rows[5].cells[0]
    p5 = c5.paragraphs[0]
    r5_title = p5.add_run("Output screenshot :\n")
    r5_title.bold = True
    r5_title.font.size = Pt(11)

    shot_items = [
        ("01_Server_Started.png", "Figure 1: Multithreaded Home Loan Server Started with Berkeley Time Daemon"),
        ("02_Cristian_Client_Sync.png", "Figure 2: Client Executing Cristian's Algorithm Clock Synchronization (SYNC_CRISTIAN)"),
        ("03_Berkeley_Daemon_Sync.png", "Figure 3: Client Receiving Berkeley Time Daemon Pulse and Applying Offset Adjustment"),
        ("04_Server_Berkeley_Daemon_Log.png", "Figure 4: Multithreaded Server Execution Log Showing Berkeley Daemon Polling & Time Adjustments")
    ]

    for fname, caption in shot_items:
        img_path = os.path.join(SCREENSHOTS_DIR, fname)
        if os.path.exists(img_path):
            p_cap = c5.add_paragraph()
            r_cap = p_cap.add_run(caption)
            r_cap.bold = True
            r_cap.font.size = Pt(9.5)
            
            p_simg = c5.add_paragraph()
            p_simg.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_simg.add_run().add_picture(img_path, width=Inches(6.2))

    # Row 6: Conclusion
    c6 = table.rows[6].cells[0]
    p6 = c6.paragraphs[0]
    r6_title = p6.add_run("Conclusion:\n")
    r6_title.bold = True
    r6_title.font.size = Pt(11)
    
    r6_body = p6.add_run(
        "In this experiment, we successfully extended our multithreaded distributed banking home loan system by implementing both Cristian's Algorithm (Client-Pull Method) and the Berkeley Algorithm (Server-Push Time Daemon). We demonstrated how Cristian's algorithm enables a client to query server time and adjust its local clock using RTT / 2 compensation. Furthermore, we showed how the Berkeley Time Daemon actively polls connected clients, computes average time skew, and broadcasts offset adjustments to correct local clock drift. This experiment effectively highlighted how clock synchronization guarantees transaction integrity in distributed banking software."
    )
    r6_body.font.size = Pt(10)

    # Row 7 & 8 Notes
    c7 = table.rows[7].cells[0]
    p7 = c7.paragraphs[0]
    r7_text = p7.add_run("Attach research paper based on experiment performed into the classroom (Not for all the practicals)")
    r7_text.font.size = Pt(9.5)
    r7_text.italic = True

    c8 = table.rows[8].cells[0]
    p8 = c8.paragraphs[0]
    r8_text = p8.add_run("(Students are required to prepare one doc file for each experiment and then upload it in pdf form.)")
    r8_text.font.size = Pt(9.5)
    r8_text.italic = True

    doc.save(DOCX_PATH)
    print("Exact SPIT Template DOCX build complete for Exp 3:", DOCX_PATH)

# -------------------------------------------------------------
# 2. BUILD EXACT SPIT TEMPLATE PDF
# -------------------------------------------------------------
def build_exact_spit_pdf():
    print("Generating Exact SPIT Template PDF report for Exp 3...")
    doc = SimpleDocTemplate(
        PDF_PATH,
        pagesize=letter,
        rightMargin=45, leftMargin=45, topMargin=45, bottomMargin=45
    )
    
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'HeaderTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        alignment=1,
        textColor=colors.HexColor("#000000")
    )
    
    meta_style = ParagraphStyle(
        'MetaText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13.5
    )

    section_heading = ParagraphStyle(
        'SecHeading',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=15,
        textColor=colors.HexColor("#000000"),
        spaceBefore=6,
        spaceAfter=4
    )

    body_style = ParagraphStyle(
        'BodyTextCustom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor("#111111")
    )

    code_style = ParagraphStyle(
        'CodeStyle',
        parent=styles['Normal'],
        fontName='Courier',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor("#002B36")
    )

    story = []

    # Header Table
    logo_img = Image(LOGO_PATH, width=70, height=70) if os.path.exists(LOGO_PATH) else Paragraph("", body_style)
    text_p = Paragraph("<b>Bhartiya Vidya Bhavans</b><br/><b>Sardar Patel Institute of Technology</b><br/>Munshi Nagar – Andheri West – 400058", title_style)
    
    head_table = Table([[logo_img, text_p]], colWidths=[90, 430])
    head_table.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('ALIGN', (0,0), (0,0), 'CENTER'),
        ('ALIGN', (1,0), (1,0), 'CENTER'),
    ]))
    story.append(head_table)
    story.append(Spacer(1, 4))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.black, spaceBefore=2, spaceAfter=6))

    # Metadata Block Table
    meta_table_data = [
        [
            Paragraph("<b>Name:</b> Atharv patil (2024300165)<br/>&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;<b>Tanish parakh (2024300157)</b><br/>&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;<b>Gitesh patil (2024300166)</b><br/><b>Branch:</b> C.E", meta_style),
            Paragraph("<b>DIVISION:</b> C<br/><b>BATCH:</b> C1<br/><b>DATE:</b> 23/08/2026", meta_style)
        ]
    ]
    t_meta = Table(meta_table_data, colWidths=[340, 180])
    t_meta.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(t_meta)
    story.append(HRFlowable(width="100%", thickness=1, color=colors.black, spaceBefore=4, spaceAfter=8))

    # Section 1: Experiment No.
    story.append(Paragraph("<b>Experiment No.: 3</b>", section_heading))
    story.append(Spacer(1, 4))

    # Section 2: Aim
    story.append(Paragraph("<b>Aim :</b> To extend a Java-based Multithreaded Distributed Banking Home Loan Management System by implementing Cristian's Algorithm (Client-Pull Method) and the Berkeley Algorithm (Server-Push Time Daemon), ensuring precise clock synchronization and accurate timestamping for concurrent loan applications and EMI transactions across distributed clients.", body_style))
    story.append(Spacer(1, 8))

    # Section 3: Theory
    story.append(Paragraph("<b>Theory :</b>", section_heading))
    story.append(Paragraph(
        "Clock Synchronization in Distributed Banking Systems is critical to align local system times across independent client nodes and servers, maintaining a globally consistent transaction ledger. In distributed banking, if physical clocks drift due to hardware oscillator variance, financial transactions (loan submissions, EMI account debits) can be logged out of order.<br/><br/>"
        "<b>1. Cristian's Algorithm (Client-Pull Method):</b><br/>"
        "• Concept: A Loan Client actively requests the exact time from the central Loan Server right before executing an EMI transaction.<br/>"
        "• Implementation: The client records local time T0 before dispatching the SYNC_CRISTIAN request to the server.<br/>"
        "• Server Response: The server receives this request and replies with its current system time T_server.<br/>"
        "• Calculation: The client receives the response at T1, computes Round Trip Time RTT = T1 - T0, and adjusts its clock to T_sync = T_server + (RTT / 2).<br/><br/>"
        "<b>2. Berkeley Algorithm (Server-Push Time Daemon Method):</b><br/>"
        "• Concept: The central Loan Server acts as an active Time Daemon, ensuring all connected banking clients are synchronized.<br/>"
        "• Implementation: A background daemon thread in the server periodically polls active client sockets with a POLL_TIME prompt.<br/>"
        "• Calculation: Clients respond with their local timestamps (TIME_REPLY). The server calculates the average time skew and sends individual offset adjustments back so clients correct local drift.<br/><br/>"
        "<b>Key Concepts Demonstrated:</b><br/>"
        "● <b>Cristian's Algorithm:</b> Adjusts client clock by querying time server and compensating for RTT / 2.<br/>"
        "● <b>Berkeley Algorithm:</b> Utilizes a central time daemon to poll clients, calculate time skew, and issue offset adjustments.<br/>"
        "● <b>Clock Synchronization:</b> Aligns local system times across distributed banking nodes.<br/>"
        "● <b>Transaction Integrity:</b> Guarantees concurrent home loan submissions and EMI payments are recorded with synchronized timestamps.",
        body_style
    ))
    story.append(Spacer(1, 10))

    # Section 4: Problem-solving & Architecture
    story.append(Paragraph("<b>Problem-solving :</b>", section_heading))
    story.append(Spacer(1, 4))
    if os.path.exists(ARCH_PATH):
        story.append(Image(ARCH_PATH, width=480, height=275))
    story.append(Spacer(1, 10))

    # Section 5: Code
    story.append(Paragraph("<b>Code :</b>", section_heading))
    story.append(Spacer(1, 4))
    story.append(Paragraph("<b>1. MultiThreadedLoanServer.java (Server with Cristian & Berkeley Daemon)</b>", ParagraphStyle('SubHead', parent=body_style, fontName='Helvetica-Bold', textColor=colors.HexColor("#006600"))))
    story.append(Spacer(1, 4))

    def create_code_paragraphs(code_text):
        lines = code_text.splitlines()
        flowables = []
        chunk = []
        for line in lines:
            chunk.append(line)
            if len(chunk) >= 40:
                chunk_str = "\n".join(chunk).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>").replace(" ", "&nbsp;")
                flowables.append(Paragraph(chunk_str, code_style))
                flowables.append(Spacer(1, 4))
                chunk = []
        if chunk:
            chunk_str = "\n".join(chunk).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>").replace(" ", "&nbsp;")
            flowables.append(Paragraph(chunk_str, code_style))
        return flowables

    story.extend(create_code_paragraphs(SERVER_CODE))
    story.append(Spacer(1, 10))

    story.append(Paragraph("<b>2. LoanClient.java (Client with Cristian Sync & Berkeley Listener)</b>", ParagraphStyle('SubHead2', parent=body_style, fontName='Helvetica-Bold', textColor=colors.HexColor("#006600"))))
    story.append(Spacer(1, 4))
    story.extend(create_code_paragraphs(CLIENT_CODE))
    story.append(Spacer(1, 12))

    # Section 6: Output screenshot
    story.append(Paragraph("<b>Output screenshot :</b>", section_heading))
    story.append(Spacer(1, 6))

    shot_items = [
        ("01_Server_Started.png", "Figure 1: Multithreaded Home Loan Server Started with Berkeley Time Daemon"),
        ("02_Cristian_Client_Sync.png", "Figure 2: Client Executing Cristian's Algorithm Clock Synchronization (SYNC_CRISTIAN)"),
        ("03_Berkeley_Daemon_Sync.png", "Figure 3: Client Receiving Berkeley Time Daemon Pulse and Applying Offset Adjustment"),
        ("04_Server_Berkeley_Daemon_Log.png", "Figure 4: Multithreaded Server Execution Log Showing Berkeley Daemon Polling & Time Adjustments")
    ]

    for fname, caption in shot_items:
        img_p = os.path.join(SCREENSHOTS_DIR, fname)
        if os.path.exists(img_p):
            story.append(Paragraph(f"<b>{caption}</b>", body_style))
            story.append(Spacer(1, 4))
            story.append(Image(img_p, width=480, height=270))
            story.append(Spacer(1, 10))

    # Section 7: Conclusion
    story.append(Paragraph("<b>Conclusion:</b> (It should include what the student has learned/ understood from the experiment).", section_heading))
    story.append(Paragraph(
        "In this experiment, we successfully extended our multithreaded distributed banking home loan system by implementing both Cristian's Algorithm (Client-Pull Method) and the Berkeley Algorithm (Server-Push Time Daemon). We demonstrated how Cristian's algorithm enables a client to query server time and adjust its local clock using RTT / 2 compensation. Furthermore, we showed how the Berkeley Time Daemon actively polls connected clients, computes average time skew, and broadcasts offset adjustments to correct local clock drift. This experiment effectively highlighted how clock synchronization guarantees transaction integrity in distributed banking software.",
        body_style
    ))
    story.append(Spacer(1, 10))

    story.append(Paragraph("<i>Attach research paper based on experiment performed into the classroom (Not for all the practicals)</i>", body_style))
    story.append(Spacer(1, 4))
    story.append(Paragraph("<i>(Students are required to prepare one doc file for each experiment and then upload it in pdf form.)</i>", body_style))

    doc.build(story, onFirstPage=draw_page_border, onLaterPages=draw_page_border)
    print("Exact SPIT Template PDF build complete for Exp 3:", PDF_PATH)

if __name__ == "__main__":
    build_exact_spit_docx()
    build_exact_spit_pdf()
