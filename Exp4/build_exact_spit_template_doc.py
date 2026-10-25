import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Image, Table, TableStyle, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

EXP4_DIR = r"c:\Users\Tanish\OneDrive\Desktop\DC Home Loan\Exp4"
PDF_PATH = os.path.join(EXP4_DIR, "DC_Lab_Exp_4_ElectionAlgo.pdf")
DOCX_PATH = os.path.join(EXP4_DIR, "DC_Lab_Exp_4_ElectionAlgo.docx")

BULLY_CODE = open(os.path.join(EXP4_DIR, "BullyElectionBanking.java"), "r", encoding="utf-8").read()

SCREENSHOTS_DIR = os.path.join(EXP4_DIR, "screenshots")
LOGO_PATH = r"c:\Users\Tanish\OneDrive\Desktop\DC Home Loan\Exp2\spit_logo.png"
ARCH_PATH = os.path.join(EXP4_DIR, "architecture_diagram.png")

def draw_page_border(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(colors.black)
    canvas.setLineWidth(1)
    canvas.rect(30, 30, doc.pagesize[0] - 60, doc.pagesize[1] - 60)
    canvas.restoreState()

# -------------------------------------------------------------
# 1. BUILD EXACT SPIT TEMPLATE DOCX
# -------------------------------------------------------------
def build_exact_spit_docx():
    print("Generating Exact SPIT Template DOCX report for Exp 4...")
    doc = docx.Document()
    
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    # Header Table: Logo + Text
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_cells = header_table.rows[0].cells
    
    p_logo = header_cells[0].paragraphs[0]
    if os.path.exists(LOGO_PATH):
        p_logo.add_run().add_picture(LOGO_PATH, width=Inches(1.2))
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    p_head = header_cells[1].paragraphs[0]
    p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p_head.add_run("Bhartiya Vidya Bhavans\n")
    r1.bold = True
    r1.font.size = Pt(13)
    r1.font.name = 'Arial'
    
    r2 = p_head.add_run("Sardar Patel Institute of Technology\n")
    r2.bold = True
    r2.font.size = Pt(13)
    r2.font.name = 'Arial'

    r3 = p_head.add_run("Munshi Nagar – Andheri West – 400058")
    r3.font.size = Pt(11)
    r3.font.name = 'Arial'

    p_div = doc.add_paragraph("—" * 65)
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Metadata Table
    meta_table = doc.add_table(rows=1, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = meta_table.rows[0].cells
    
    hdr_cells[0].paragraphs[0].add_run("Name: Atharv patil (UID: 2024300165)\n          Tanish parakh (UID: 2024300157)\n          Gitesh patil (UID: 2024300166)\nBranch: C.E").bold = True
    hdr_cells[1].paragraphs[0].add_run("DIVISION: C\nBATCH: C1\nDATE: 24/08/2026").bold = True
    
    p_div2 = doc.add_paragraph("—" * 65)
    p_div2.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Main Template Table with specific section rows
    table = doc.add_table(rows=9, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Row 0: Experiment No.
    c0 = table.rows[0].cells[0]
    p0 = c0.paragraphs[0]
    r0 = p0.add_run("Experiment No.: 4")
    r0.bold = True
    r0.font.size = Pt(11)

    # Row 1: Aim
    c1 = table.rows[1].cells[0]
    p1 = c1.paragraphs[0]
    r1_title = p1.add_run("Aim : ")
    r1_title.bold = True
    r1_title.font.size = Pt(11)
    r1_body = p1.add_run("To implement and evaluate The Bully Algorithm for Leader / Coordinator Election in a Distributed Banking Home Loan Management System, demonstrating dynamic election of the highest active process ID as the Central Master Coordinator upon coordinator failure or recovery.")
    r1_body.font.size = Pt(10)

    # Row 2: Theory
    c2 = table.rows[2].cells[0]
    p2 = c2.paragraphs[0]
    r2_title = p2.add_run("Theory :\n")
    r2_title.bold = True
    r2_title.font.size = Pt(11)
    r2_body = p2.add_run(
        "In a distributed system, a Coordinator Node (Leader) is responsible for central management, centralized locking, transaction approval, and resource allocation. If the coordinator process crashes or fails, the system must dynamically run an Election Algorithm to choose a new coordinator among active processes.\n\n"
        "1. The Bully Algorithm (Garcia-Molina):\n"
        "The Bully Algorithm assumes that every process has a unique, globally known Process ID (priority). The process with the highest active Process ID is always elected as the Coordinator ('bullying' lower ID processes).\n"
        "Algorithm Steps:\n"
        "• Election Initiation: When process P_i detects that the current coordinator is unresponsive, it sends an ELECTION message to all processes with higher IDs (P_{i+1}, ..., P_N).\n"
        "• Response (OK / ANSWER): If a higher process P_k is alive, it sends an OK message back to P_i and starts its own election.\n"
        "• Timeout & Victory: If no higher process responds within a timeout period T, process P_i declares itself the winner and broadcasts a COORDINATOR message to all lower active processes.\n"
        "• Process Recovery: If a crashed process with a higher ID comes back online, it immediately initiates an election and reclaims the coordinator role.\n\n"
        "2. Message Complexity Analysis:\n"
        "• Best-Case Complexity: O(N) — When the process with the second-highest ID detects failure.\n"
        "• Worst-Case Complexity: O(N^2) — When the process with the lowest ID detects failure.\n\n"
        "Key Concepts Demonstrated:\n"
        "● Leader Election: Dynamic coordinator selection without human intervention.\n"
        "● High Availability: Guarantees distributed banking system continuity upon node crashes.\n"
        "● Bully Behavior: Highest priority active process always reclaims control."
    )
    r2_body.font.size = Pt(10)

    # Row 3: Problem-solving
    c3 = table.rows[3].cells[0]
    p3 = c3.paragraphs[0]
    r3_title = p3.add_run("Problem-solving :\n")
    r3_title.bold = True
    r3_title.font.size = Pt(11)
    
    if os.path.exists(ARCH_PATH):
        p3_img = c3.add_paragraph()
        p3_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p3_img.add_run().add_picture(ARCH_PATH, width=Inches(6.2))

    # Row 4: Code
    c4 = table.rows[4].cells[0]
    p4 = c4.paragraphs[0]
    r4_title = p4.add_run("Code :\n")
    r4_title.bold = True
    r4_title.font.size = Pt(11)

    r4_c1 = p4.add_run("1. BullyElectionBanking.java (The Bully Algorithm Implementation)\n")
    r4_c1.bold = True
    r4_c1.font.color.rgb = RGBColor(0, 102, 0)

    p4_code1 = c4.add_paragraph()
    r4_code1_text = p4_code1.add_run(BULLY_CODE)
    r4_code1_text.font.name = 'Consolas'
    r4_code1_text.font.size = Pt(8.5)

    # Row 5: Output screenshot
    c5 = table.rows[5].cells[0]
    p5 = c5.paragraphs[0]
    r5_title = p5.add_run("Output screenshot :\n")
    r5_title.bold = True
    r5_title.font.size = Pt(11)

    shot_items = [
        ("01_Initial_Cluster_State.png", "Figure 1: Initial Cluster State with BankNode-4 as Active Coordinator"),
        ("02_Coordinator_Crash_Event.png", "Figure 2: Central Master Coordinator Crash Event & Election Initiation by BankNode-2"),
        ("03_Election_Initiation_P3_Wins.png", "Figure 3: BankNode-3 Winning Election & Broadcasting COORDINATOR Announcement"),
        ("04_Higher_Node_Recovery_Bully.png", "Figure 4: BankNode-4 Recovery & Reclaiming Leadership as Central Master Coordinator")
    ]

    for fname, caption in shot_items:
        img_path = os.path.join(SCREENSHOTS_DIR, fname)
        if os.path.exists(img_path):
            p_cap = c5.add_paragraph()
            r_cap = p_cap.add_run(caption)
            r_cap.bold = True
            r_cap.font.size = Pt(9.5)
            
            p_simg = c5.add_paragraph()
            p_simg.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_simg.add_run().add_picture(img_path, width=Inches(6.2))

    # Row 6: Conclusion
    c6 = table.rows[6].cells[0]
    p6 = c6.paragraphs[0]
    r6_title = p6.add_run("Conclusion:\n")
    r6_title.bold = True
    r6_title.font.size = Pt(11)
    
    r6_body = p6.add_run(
        "In this experiment, we successfully implemented and analyzed The Bully Algorithm for Leader / Coordinator Election in a Distributed Banking System. We demonstrated how active banking nodes handle coordinator failure by initiating election messages, exchanging OK responses, and electing the highest active Process ID (BankNode-3) as the new master. Furthermore, we showed how a recovered node (BankNode-4) initiates an election to reclaim its leadership position. This experiment effectively highlighted the fault-tolerance and high-availability properties of election algorithms in distributed systems."
    )
    r6_body.font.size = Pt(10)

    # Row 7 & 8 Notes
    c7 = table.rows[7].cells[0]
    p7 = c7.paragraphs[0]
    r7_text = p7.add_run("Attach research paper based on experiment performed into the classroom (Not for all the practicals)")
    r7_text.font.size = Pt(9.5)
    r7_text.italic = True

    c8 = table.rows[8].cells[0]
    p8 = c8.paragraphs[0]
    r8_text = p8.add_run("(Students are required to prepare one doc file for each experiment and then upload it in pdf form.)")
    r8_text.font.size = Pt(9.5)
    r8_text.italic = True

    doc.save(DOCX_PATH)
    print("Exact SPIT Template DOCX build complete for Exp 4:", DOCX_PATH)

# -------------------------------------------------------------
# 2. BUILD EXACT SPIT TEMPLATE PDF
# -------------------------------------------------------------
def build_exact_spit_pdf():
    print("Generating Exact SPIT Template PDF report for Exp 4...")
    doc = SimpleDocTemplate(
        PDF_PATH,
        pagesize=letter,
        rightMargin=45, leftMargin=45, topMargin=45, bottomMargin=45
    )
    
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'HeaderTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        alignment=1,
        textColor=colors.HexColor("#000000")
    )
    
    meta_style = ParagraphStyle(
        'MetaText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13.5
    )

    section_heading = ParagraphStyle(
        'SecHeading',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=15,
        textColor=colors.HexColor("#000000"),
        spaceBefore=6,
        spaceAfter=4
    )

    body_style = ParagraphStyle(
        'BodyTextCustom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor("#111111")
    )

    code_style = ParagraphStyle(
        'CodeStyle',
        parent=styles['Normal'],
        fontName='Courier',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor("#002B36")
    )

    story = []

    # Header Table
    logo_img = Image(LOGO_PATH, width=70, height=70) if os.path.exists(LOGO_PATH) else Paragraph("", body_style)
    text_p = Paragraph("<b>Bhartiya Vidya Bhavans</b><br/><b>Sardar Patel Institute of Technology</b><br/>Munshi Nagar – Andheri West – 400058", title_style)
    
    head_table = Table([[logo_img, text_p]], colWidths=[90, 430])
    head_table.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('ALIGN', (0,0), (0,0), 'CENTER'),
        ('ALIGN', (1,0), (1,0), 'CENTER'),
    ]))
    story.append(head_table)
    story.append(Spacer(1, 4))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.black, spaceBefore=2, spaceAfter=6))

    # Metadata Block Table
    meta_table_data = [
        [
            Paragraph("<b>Name:</b> Atharv patil (2024300165)<br/>&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;<b>Tanish parakh (2024300157)</b><br/>&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;<b>Gitesh patil (2024300166)</b><br/><b>Branch:</b> C.E", meta_style),
            Paragraph("<b>DIVISION:</b> C<br/><b>BATCH:</b> C1<br/><b>DATE:</b> 24/08/2026", meta_style)
        ]
    ]
    t_meta = Table(meta_table_data, colWidths=[340, 180])
    t_meta.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(t_meta)
    story.append(HRFlowable(width="100%", thickness=1, color=colors.black, spaceBefore=4, spaceAfter=8))

    # Section 1: Experiment No.
    story.append(Paragraph("<b>Experiment No.: 4</b>", section_heading))
    story.append(Spacer(1, 4))

    # Section 2: Aim
    story.append(Paragraph("<b>Aim :</b> To implement and evaluate The Bully Algorithm for Leader / Coordinator Election in a Distributed Banking Home Loan Management System, demonstrating dynamic election of the highest active process ID as the Central Master Coordinator upon coordinator failure or recovery.", body_style))
    story.append(Spacer(1, 8))

    # Section 3: Theory
    story.append(Paragraph("<b>Theory :</b>", section_heading))
    story.append(Paragraph(
        "In a distributed system, a Coordinator Node (Leader) is responsible for central management, centralized locking, transaction approval, and resource allocation. If the coordinator process crashes or fails, the system must dynamically run an Election Algorithm to choose a new coordinator among active processes.<br/><br/>"
        "<b>1. The Bully Algorithm (Garcia-Molina):</b><br/>"
        "The Bully Algorithm assumes that every process has a unique, globally known Process ID (priority). The process with the highest active Process ID is always elected as the Coordinator ('bullying' lower ID processes).<br/>"
        "Algorithm Steps:<br/>"
        "• <b>Election Initiation:</b> When process P_i detects that the current coordinator is unresponsive, it sends an ELECTION message to all processes with higher IDs (P_{i+1}, ..., P_N).<br/>"
        "• <b>Response (OK / ANSWER):</b> If a higher process P_k is alive, it sends an OK message back to P_i and starts its own election.<br/>"
        "• <b>Timeout & Victory:</b> If no higher process responds within a timeout period T, process P_i declares itself the winner and broadcasts a COORDINATOR message to all lower active processes.<br/>"
        "• <b>Process Recovery:</b> If a crashed process with a higher ID comes back online, it immediately initiates an election and reclaims the coordinator role.<br/><br/>"
        "<b>2. Message Complexity Analysis:</b><br/>"
        "• <b>Best-Case Complexity:</b> O(N) — When the process with the second-highest ID detects failure.<br/>"
        "• <b>Worst-Case Complexity:</b> O(N^2) — When the process with the lowest ID detects failure.<br/><br/>"
        "<b>Key Concepts Demonstrated:</b><br/>"
        "● <b>Leader Election:</b> Dynamic coordinator selection without human intervention.<br/>"
        "● <b>High Availability:</b> Guarantees distributed banking system continuity upon node crashes.<br/>"
        "● <b>Bully Behavior:</b> Highest priority active process always reclaims control.",
        body_style
    ))
    story.append(Spacer(1, 10))

    # Section 4: Problem-solving & Architecture
    story.append(Paragraph("<b>Problem-solving :</b>", section_heading))
    story.append(Spacer(1, 4))
    if os.path.exists(ARCH_PATH):
        story.append(Image(ARCH_PATH, width=480, height=275))
    story.append(Spacer(1, 10))

    # Section 5: Code
    story.append(Paragraph("<b>Code :</b>", section_heading))
    story.append(Spacer(1, 4))
    story.append(Paragraph("<b>1. BullyElectionBanking.java (The Bully Algorithm Implementation)</b>", ParagraphStyle('SubHead', parent=body_style, fontName='Helvetica-Bold', textColor=colors.HexColor("#006600"))))
    story.append(Spacer(1, 4))

    def create_code_paragraphs(code_text):
        lines = code_text.splitlines()
        flowables = []
        chunk = []
        for line in lines:
            chunk.append(line)
            if len(chunk) >= 40:
                chunk_str = "\n".join(chunk).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>").replace(" ", "&nbsp;")
                flowables.append(Paragraph(chunk_str, code_style))
                flowables.append(Spacer(1, 4))
                chunk = []
        if chunk:
            chunk_str = "\n".join(chunk).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>").replace(" ", "&nbsp;")
            flowables.append(Paragraph(chunk_str, code_style))
        return flowables

    story.extend(create_code_paragraphs(BULLY_CODE))
    story.append(Spacer(1, 12))

    # Section 6: Output screenshot
    story.append(Paragraph("<b>Output screenshot :</b>", section_heading))
    story.append(Spacer(1, 6))

    shot_items = [
        ("01_Initial_Cluster_State.png", "Figure 1: Initial Cluster State with BankNode-4 as Active Coordinator"),
        ("02_Coordinator_Crash_Event.png", "Figure 2: Central Master Coordinator Crash Event & Election Initiation by BankNode-2"),
        ("03_Election_Initiation_P3_Wins.png", "Figure 3: BankNode-3 Winning Election & Broadcasting COORDINATOR Announcement"),
        ("04_Higher_Node_Recovery_Bully.png", "Figure 4: BankNode-4 Recovery & Reclaiming Leadership as Central Master Coordinator")
    ]

    for fname, caption in shot_items:
        img_p = os.path.join(SCREENSHOTS_DIR, fname)
        if os.path.exists(img_p):
            story.append(Paragraph(f"<b>{caption}</b>", body_style))
            story.append(Spacer(1, 4))
            story.append(Image(img_p, width=480, height=270))
            story.append(Spacer(1, 10))

    # Section 7: Conclusion
    story.append(Paragraph("<b>Conclusion:</b> (It should include what the student has learned/ understood from the experiment).", section_heading))
    story.append(Paragraph(
        "In this experiment, we successfully implemented and analyzed The Bully Algorithm for Leader / Coordinator Election in a Distributed Banking System. We demonstrated how active banking nodes handle coordinator failure by initiating election messages, exchanging OK responses, and electing the highest active Process ID (BankNode-3) as the new master. Furthermore, we showed how a recovered node (BankNode-4) initiates an election to reclaim its leadership position. This experiment effectively highlighted the fault-tolerance and high-availability properties of election algorithms in distributed systems.",
        body_style
    ))
    story.append(Spacer(1, 10))

    story.append(Paragraph("<i>Attach research paper based on experiment performed into the classroom (Not for all the practicals)</i>", body_style))
    story.append(Spacer(1, 4))
    story.append(Paragraph("<i>(Students are required to prepare one doc file for each experiment and then upload it in pdf form.)</i>", body_style))

    doc.build(story, onFirstPage=draw_page_border, onLaterPages=draw_page_border)
    print("Exact SPIT Template PDF build complete for Exp 4:", PDF_PATH)

if __name__ == "__main__":
    build_exact_spit_docx()
    build_exact_spit_pdf()
