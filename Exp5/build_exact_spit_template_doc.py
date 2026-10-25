import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Image, Table, TableStyle, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

EXP5_DIR = r"c:\Users\Tanish\OneDrive\Desktop\DC Home Loan\Exp5"
PDF_PATH = os.path.join(EXP5_DIR, "DC_Lab_Exp_5_FaultTolerance.pdf")
DOCX_PATH = os.path.join(EXP5_DIR, "DC_Lab_Exp_5_FaultTolerance.docx")

SERVER_CODE = open(os.path.join(EXP5_DIR, "PrimaryBackupLoanServer.java"), "r", encoding="utf-8").read()
CLIENT_CODE = open(os.path.join(EXP5_DIR, "FailoverLoanClient.java"), "r", encoding="utf-8").read()

SCREENSHOTS_DIR = os.path.join(EXP5_DIR, "screenshots")
LOGO_PATH = r"c:\Users\Tanish\OneDrive\Desktop\DC Home Loan\Exp2\spit_logo.png"
ARCH_PATH = os.path.join(EXP5_DIR, "architecture_diagram.png")

def draw_page_border(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(colors.black)
    canvas.setLineWidth(1)
    canvas.rect(30, 30, doc.pagesize[0] - 60, doc.pagesize[1] - 60)
    canvas.restoreState()

# -------------------------------------------------------------
# 1. BUILD EXACT SPIT TEMPLATE DOCX
# -------------------------------------------------------------
def build_exact_spit_docx():
    print("Generating Exact SPIT Template DOCX report for Exp 5...")
    doc = docx.Document()
    
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    # Header Table: Logo + Text
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_cells = header_table.rows[0].cells
    
    p_logo = header_cells[0].paragraphs[0]
    if os.path.exists(LOGO_PATH):
        p_logo.add_run().add_picture(LOGO_PATH, width=Inches(1.2))
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    p_head = header_cells[1].paragraphs[0]
    p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p_head.add_run("Bhartiya Vidya Bhavans\n")
    r1.bold = True
    r1.font.size = Pt(13)
    r1.font.name = 'Arial'
    
    r2 = p_head.add_run("Sardar Patel Institute of Technology\n")
    r2.bold = True
    r2.font.size = Pt(13)
    r2.font.name = 'Arial'

    r3 = p_head.add_run("Munshi Nagar – Andheri West – 400058")
    r3.font.size = Pt(11)
    r3.font.name = 'Arial'

    p_div = doc.add_paragraph("—" * 65)
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Metadata Table
    meta_table = doc.add_table(rows=1, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = meta_table.rows[0].cells
    
    hdr_cells[0].paragraphs[0].add_run("Name: Atharv Patil (UID: 2024300165)\n          Tanish Parakh (UID: 2024300157)\n          Gitesh Patil (UID: 2024300166)\nBranch: C.E").bold = True
    hdr_cells[1].paragraphs[0].add_run("DIVISION: C\nBATCH: C1\nDATE: 28/09/2026").bold = True
    
    p_div2 = doc.add_paragraph("—" * 65)
    p_div2.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Main Template Table with specific section rows
    table = doc.add_table(rows=9, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Row 0: Experiment No.
    c0 = table.rows[0].cells[0]
    p0 = c0.paragraphs[0]
    r0 = p0.add_run("Experiment No.: 5")
    r0.bold = True
    r0.font.size = Pt(11)

    # Row 1: Aim
    c1 = table.rows[1].cells[0]
    p1 = c1.paragraphs[0]
    r1_title = p1.add_run("Aim : ")
    r1_title.bold = True
    r1_title.font.size = Pt(11)
    r1_body = p1.add_run("To implement Fault Tolerance in a Java-based Distributed Banking Home Loan Management System using Primary-Backup (Passive) Replication, ensuring that the loan database survives a coordinator/server crash without loss of data and without service interruption to connected clients.")
    r1_body.font.size = Pt(10)

    # Row 2: Theory
    c2 = table.rows[2].cells[0]
    p2 = c2.paragraphs[0]
    r2_title = p2.add_run("Theory :\n")
    r2_title.bold = True
    r2_title.font.size = Pt(11)
    r2_body = p2.add_run(
        "Fault Tolerance is the property of a distributed system that allows it to continue operating correctly even when one or more of its components fail. In a banking system, the loan server cannot be a single point of failure — if the coordinator (Master, as elected via the Bully Algorithm in Experiment 4) crashes, in-flight loan submissions and EMI records must not be lost, and clients must still be able to reach a working server.\n\n"
        "Replication is the standard technique used to achieve fault tolerance: multiple copies (replicas) of the server state (the loanDatabase) are maintained on different nodes so that if one node fails, another already has an up-to-date copy and can take over.\n\n"
        "1. Replication Strategies:\n"
        "• Active Replication: Every replica independently processes every client request in parallel. Requires strict request ordering (e.g., via Lamport clocks) and is expensive to keep deterministic.\n"
        "• Passive / Primary-Backup Replication: Only one replica (the Primary) processes client requests. After each state-changing operation, the Primary forwards (replicates) the update to one or more Backup replicas. Backups stay passive until the Primary fails, at which point a Backup is promoted to become the new Primary. This experiment implements Primary-Backup replication.\n\n"
        "2. Primary-Backup Replication — Mechanism:\n"
        "• Primary Role: Accepts all client connections on the well-known client port, processes SUBMIT / PAY_EMI / STATUS requests, updates its local loanDatabase, and returns a response to the client.\n"
        "• Write Propagation: Immediately after a write (SUBMIT / PAY_EMI), the Primary sends a REPLICATE message containing the same request over a dedicated replication socket to the Backup, which applies the identical update to its own copy of loanDatabase — keeping the two copies consistent.\n"
        "• Heartbeat / Failure Detection: The Primary periodically sends a PING message to the Backup over the replication channel. The Backup runs a monitor thread that checks the time elapsed since the last PING. If this exceeds a configured timeout, the Primary is presumed to have crashed.\n"
        "• Failover / Promotion: On detecting a timeout, the Backup promotes itself — it stops passively listening only for replication traffic and opens its own client-facing ServerSocket, becoming the new Primary. Because every prior write was already replicated and applied before being acknowledged, the promoted node has zero data loss.\n"
        "• Client-Side Failover: The LoanClient is made resilient by first attempting to connect to the known Primary address; if the connection is refused (socket exception), it retries against the failover target (the former Backup, now Primary), so the customer-facing service keeps working transparently.\n\n"
        "3. Consistency Guarantee:\n"
        "Because the Backup only needs to apply a REPLICATE message locally, the promoted node's loanDatabase reflects every write that was ever confirmed to a client. This gives the system the fault-tolerance property: continued correct operation despite a single-node crash.\n\n"
        "Real-Life Examples:\n"
        "• MySQL / PostgreSQL master-replica setups: A primary database handles writes; replicas apply the replication log and can be promoted on primary failure.\n"
        "• Redis Sentinel: Monitors a Redis master and automatically promotes a replica if the master becomes unreachable.\n"
        "• Banking core systems: A standby data-centre replica takes over transaction processing if the primary data-centre goes offline.\n\n"
        "Key Concepts Demonstrated:\n"
        "● Primary: Single node that owns all writes and serves clients.\n"
        "● Backup: Passive replica that mirrors every write and watches for Primary failure.\n"
        "● Replication: Propagation of state-changing requests from Primary to Backup.\n"
        "● Heartbeat Monitoring: Periodic liveness signal used to detect crashes.\n"
        "● Failover: Automatic promotion of the Backup to Primary with no data loss.\n"
        "● Client Resilience: Automatic client-side reconnection to the newly promoted Primary."
    )
    r2_body.font.size = Pt(10)

    # Row 3: Problem-solving
    c3 = table.rows[3].cells[0]
    p3 = c3.paragraphs[0]
    r3_title = p3.add_run("Problem-solving :\n")
    r3_title.bold = True
    r3_title.font.size = Pt(11)
    
    r3_desc = p3.add_run(
        "• The single Primary server (Exp 1-4) is a single point of failure — if it crashes mid-transaction, all clients lose service.\n"
        "• Solution: Run a second node (Backup) that continuously mirrors the Primary's loan database via replication messages.\n"
        "• A replication channel (separate port) carries write requests (REPLICATE|...) and periodic heartbeats (PING).\n"
        "• The Backup runs a heartbeat-monitor thread; if no PING arrives within the timeout window, it assumes the Primary is down.\n"
        "• On timeout, the Backup promotes itself: it starts its own client-facing ServerSocket and begins serving requests using the already-replicated loanDatabase.\n"
        "• The LoanClient adds a failover step: if it cannot connect to the Primary's address, it retries the Backup's address so the customer experience is uninterrupted.\n\n"
    )
    r3_desc.font.size = Pt(10)
    
    if os.path.exists(ARCH_PATH):
        p3_img = c3.add_paragraph()
        p3_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p3_img.add_run().add_picture(ARCH_PATH, width=Inches(6.2))

    # Row 4: Code
    c4 = table.rows[4].cells[0]
    p4 = c4.paragraphs[0]
    r4_title = p4.add_run("Code :\n")
    r4_title.bold = True
    r4_title.font.size = Pt(11)

    r4_c1 = p4.add_run("1. PrimaryBackupLoanServer.java (extends the Exp 4 BankNode-style server with replication + failover)\n")
    r4_c1.bold = True
    r4_c1.font.color.rgb = RGBColor(0, 102, 0)

    p4_code1 = c4.add_paragraph()
    r4_code1_text = p4_code1.add_run(SERVER_CODE)
    r4_code1_text.font.name = 'Consolas'
    r4_code1_text.font.size = Pt(8.5)

    p4_c2 = c4.add_paragraph()
    r4_c2_t = p4_c2.add_run("\n2. FailoverLoanClient.java (extends the Exp 1/2 LoanClient with automatic reconnect)\n")
    r4_c2_t.bold = True
    r4_c2_t.font.color.rgb = RGBColor(0, 102, 0)

    p4_code2 = c4.add_paragraph()
    r4_code2_text = p4_code2.add_run(CLIENT_CODE)
    r4_code2_text.font.name = 'Consolas'
    r4_code2_text.font.size = Pt(8.5)

    # Row 5: Output screenshot
    c5 = table.rows[5].cells[0]
    p5 = c5.paragraphs[0]
    r5_title = p5.add_run("Output screenshot :\n")
    r5_title.bold = True
    r5_title.font.size = Pt(11)

    shot_items = [
        # Real Java Swing Frontend Screenshots
        ("01_Real_Java_UI_Dashboard.png", "Figure 1: BANKCORE Distributed Banking Dashboard (Java Swing) -- Real-time Node Health (Bank Node 01 Primary, Bank Node 02 Backup) and System Metrics"),
        ("02_Real_Java_UI_Loans_Management.png", "Figure 2: BANKCORE Loan Applications Interface (Java Swing) -- Active Loan Registry displaying HL1001 for Gitesh Patil (Rs 62,00,000)"),
        ("03_Real_Java_UI_Loan_Details.png", "Figure 3: BANKCORE Detailed Loan View (Java Swing) -- Verified Application Specification for HL1001 (Customer CUS201 Gitesh Patil)"),
        ("04_Real_Java_UI_Failover_Active.png", "Figure 4: BANKCORE Replication & Failover Monitor (Java Swing) -- Primary Node 01 Crash Simulation, Backup Node 02 Promotion, and Live Sync Stream"),
        
        # Real Terminal Execution Screenshots
        ("01_Primary_Backup_Startup.png", "Figure 5: Primary and Backup Nodes Starting Up (Heartbeat Monitoring Begins on Ports 5000 & 6000)"),
        ("02_Client_Submit_Replicated.png", "Figure 6: Client SUBMIT Request Processed by Primary and Replicated to Backup (Gitesh Patil, Rs 62,00,000)"),
        ("03_Primary_Crash_Failover_Promotion.png", "Figure 7: Primary Crash Simulated -- Backup Detects Timeout (>6000ms) and Self-Promotes to Primary"),
        ("04_Client_Failover_Reconnect.png", "Figure 8: Client Automatically Fails Over to the Newly Promoted Primary (Loan Data HL1001 Intact)")
    ]

    for fname, caption in shot_items:
        img_path = os.path.join(SCREENSHOTS_DIR, fname)
        if os.path.exists(img_path):
            p_cap = c5.add_paragraph()
            r_cap = p_cap.add_run(caption)
            r_cap.bold = True
            r_cap.font.size = Pt(9.5)
            
            p_simg = c5.add_paragraph()
            p_simg.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_simg.add_run().add_picture(img_path, width=Inches(6.2))

    # Row 6: Conclusion
    c6 = table.rows[6].cells[0]
    p6 = c6.paragraphs[0]
    r6_title = p6.add_run("Conclusion:\n")
    r6_title.bold = True
    r6_title.font.size = Pt(11)
    
    r6_body = p6.add_run(
        "In this experiment, we successfully implemented Fault Tolerance in our distributed banking home loan management system using Primary-Backup (Passive) Replication, extending the multithreaded, clock-synchronized, election-capable server built in Experiments 1-4. We demonstrated how the Primary node replicates every state-changing loan request to a Backup node over a dedicated replication channel, how a heartbeat mechanism allows the Backup to detect Primary failure, and how the Backup seamlessly promotes itself to the new Primary without losing any previously acknowledged loan data. We also extended the LoanClient with automatic failover so that customers experience uninterrupted service across a server crash. This experiment highlighted how replication and failure detection together provide high availability and data durability in a distributed banking system."
    )
    r6_body.font.size = Pt(10)

    # Row 7 & 8 Notes
    c7 = table.rows[7].cells[0]
    p7 = c7.paragraphs[0]
    r7_text = p7.add_run("Attach research paper based on experiment performed into the classroom (Not for all the practicals)")
    r7_text.font.size = Pt(9.5)
    r7_text.italic = True

    c8 = table.rows[8].cells[0]
    p8 = c8.paragraphs[0]
    r8_text = p8.add_run("(Students are required to prepare one doc file for each experiment and then upload it in pdf form.)")
    r8_text.font.size = Pt(9.5)
    r8_text.italic = True

    doc.save(DOCX_PATH)
    print("Exact SPIT Template DOCX build complete for Exp 5:", DOCX_PATH)

# -------------------------------------------------------------
# 2. BUILD EXACT SPIT TEMPLATE PDF
# -------------------------------------------------------------
def build_exact_spit_pdf():
    print("Generating Exact SPIT Template PDF report for Exp 5...")
    doc = SimpleDocTemplate(
        PDF_PATH,
        pagesize=letter,
        rightMargin=45, leftMargin=45, topMargin=45, bottomMargin=45
    )
    
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'HeaderTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        alignment=1,
        textColor=colors.HexColor("#000000")
    )
    
    meta_style = ParagraphStyle(
        'MetaText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13.5
    )

    section_heading = ParagraphStyle(
        'SecHeading',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=15,
        textColor=colors.HexColor("#000000"),
        spaceBefore=6,
        spaceAfter=4
    )

    body_style = ParagraphStyle(
        'BodyTextCustom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor("#111111")
    )

    code_style = ParagraphStyle(
        'CodeStyle',
        parent=styles['Normal'],
        fontName='Courier',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor("#002B36")
    )

    story = []

    # Header Table
    logo_img = Image(LOGO_PATH, width=70, height=70) if os.path.exists(LOGO_PATH) else Paragraph("", body_style)
    text_p = Paragraph("<b>Bhartiya Vidya Bhavans</b><br/><b>Sardar Patel Institute of Technology</b><br/>Munshi Nagar – Andheri West – 400058", title_style)
    
    head_table = Table([[logo_img, text_p]], colWidths=[90, 430])
    head_table.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('ALIGN', (0,0), (0,0), 'CENTER'),
        ('ALIGN', (1,0), (1,0), 'CENTER'),
    ]))
    story.append(head_table)
    story.append(Spacer(1, 4))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.black, spaceBefore=2, spaceAfter=6))

    # Metadata Block Table
    meta_table_data = [
        [
            Paragraph("<b>Name:</b> Atharv Patil (2024300165)<br/>&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;<b>Tanish Parakh (2024300157)</b><br/>&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;<b>Gitesh Patil (2024300166)</b><br/><b>Branch:</b> C.E", meta_style),
            Paragraph("<b>DIVISION:</b> C<br/><b>BATCH:</b> C1<br/><b>DATE:</b> 28/09/2026", meta_style)
        ]
    ]
    t_meta = Table(meta_table_data, colWidths=[340, 180])
    t_meta.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(t_meta)
    story.append(HRFlowable(width="100%", thickness=1, color=colors.black, spaceBefore=4, spaceAfter=8))

    # Section 1: Experiment No.
    story.append(Paragraph("<b>Experiment No.: 5</b>", section_heading))
    story.append(Spacer(1, 4))

    # Section 2: Aim
    story.append(Paragraph("<b>Aim :</b> To implement Fault Tolerance in a Java-based Distributed Banking Home Loan Management System using Primary-Backup (Passive) Replication, ensuring that the loan database survives a coordinator/server crash without loss of data and without service interruption to connected clients.", body_style))
    story.append(Spacer(1, 8))

    # Section 3: Theory
    story.append(Paragraph("<b>Theory :</b>", section_heading))
    story.append(Paragraph(
        "Fault Tolerance is the property of a distributed system that allows it to continue operating correctly even when one or more of its components fail. In a banking system, the loan server cannot be a single point of failure — if the coordinator (Master, as elected via the Bully Algorithm in Experiment 4) crashes, in-flight loan submissions and EMI records must not be lost, and clients must still be able to reach a working server.<br/><br/>"
        "Replication is the standard technique used to achieve fault tolerance: multiple copies (replicas) of the server state (the loanDatabase) are maintained on different nodes so that if one node fails, another already has an up-to-date copy and can take over.<br/><br/>"
        "<b>1. Replication Strategies:</b><br/>"
        "• <b>Active Replication:</b> Every replica independently processes every client request in parallel. Requires strict request ordering (e.g., via Lamport clocks) and is expensive to keep deterministic.<br/>"
        "• <b>Passive / Primary-Backup Replication:</b> Only one replica (the Primary) processes client requests. After each state-changing operation, the Primary forwards (replicates) the update to one or more Backup replicas. Backups stay passive until the Primary fails, at which point a Backup is promoted to become the new Primary. This experiment implements Primary-Backup replication.<br/><br/>"
        "<b>2. Primary-Backup Replication — Mechanism:</b><br/>"
        "• <b>Primary Role:</b> Accepts all client connections on the well-known client port, processes SUBMIT / PAY_EMI / STATUS requests, updates its local loanDatabase, and returns a response to the client.<br/>"
        "• <b>Write Propagation:</b> Immediately after a write (SUBMIT / PAY_EMI), the Primary sends a REPLICATE message containing the same request over a dedicated replication socket to the Backup, which applies the identical update to its own copy of loanDatabase — keeping the two copies consistent.<br/>"
        "• <b>Heartbeat / Failure Detection:</b> The Primary periodically sends a PING message to the Backup over the replication channel. The Backup runs a monitor thread that checks the time elapsed since the last PING. If this exceeds a configured timeout, the Primary is presumed to have crashed.<br/>"
        "• <b>Failover / Promotion:</b> On detecting a timeout, the Backup promotes itself — it stops passively listening only for replication traffic and opens its own client-facing ServerSocket, becoming the new Primary. Because every prior write was already replicated and applied before being acknowledged, the promoted node has zero data loss.<br/>"
        "• <b>Client-Side Failover:</b> The LoanClient is made resilient by first attempting to connect to the known Primary address; if the connection is refused (socket exception), it retries against the failover target (the former Backup, now Primary), so the customer-facing service keeps working transparently.<br/><br/>"
        "<b>3. Consistency Guarantee:</b><br/>"
        "Because the Backup only needs to apply a REPLICATE message locally, the promoted node's loanDatabase reflects every write that was ever confirmed to a client. This gives the system the fault-tolerance property: continued correct operation despite a single-node crash.<br/><br/>"
        "<b>Real-Life Examples:</b><br/>"
        "• MySQL / PostgreSQL master-replica setups: A primary database handles writes; replicas apply the replication log and can be promoted on primary failure.<br/>"
        "• Redis Sentinel: Monitors a Redis master and automatically promotes a replica if the master becomes unreachable.<br/>"
        "• Banking core systems: A standby data-centre replica takes over transaction processing if the primary data-centre goes offline.<br/><br/>"
        "<b>Key Concepts Demonstrated:</b><br/>"
        "● <b>Primary:</b> Single node that owns all writes and serves clients.<br/>"
        "● <b>Backup:</b> Passive replica that mirrors every write and watches for Primary failure.<br/>"
        "● <b>Replication:</b> Propagation of state-changing requests from Primary to Backup.<br/>"
        "● <b>Heartbeat Monitoring:</b> Periodic liveness signal used to detect crashes.<br/>"
        "● <b>Failover:</b> Automatic promotion of the Backup to Primary with no data loss.<br/>"
        "● <b>Client Resilience:</b> Automatic client-side reconnection to the newly promoted Primary.",
        body_style
    ))
    story.append(Spacer(1, 10))

    # Section 4: Problem-solving & Architecture
    story.append(Paragraph("<b>Problem-solving :</b>", section_heading))
    story.append(Spacer(1, 4))
    story.append(Paragraph(
        "• The single Primary server (Exp 1-4) is a single point of failure — if it crashes mid-transaction, all clients lose service.<br/>"
        "• Solution: Run a second node (Backup) that continuously mirrors the Primary's loan database via replication messages.<br/>"
        "• A replication channel (separate port) carries write requests (REPLICATE|...) and periodic heartbeats (PING).<br/>"
        "• The Backup runs a heartbeat-monitor thread; if no PING arrives within the timeout window, it assumes the Primary is down.<br/>"
        "• On timeout, the Backup promotes itself: it starts its own client-facing ServerSocket and begins serving requests using the already-replicated loanDatabase.<br/>"
        "• The LoanClient adds a failover step: if it cannot connect to the Primary's address, it retries the Backup's address so the customer experience is uninterrupted.<br/>",
        body_style
    ))
    story.append(Spacer(1, 6))
    if os.path.exists(ARCH_PATH):
        story.append(Image(ARCH_PATH, width=480, height=275))
    story.append(Spacer(1, 10))

    # Section 5: Code
    story.append(Paragraph("<b>Code :</b>", section_heading))
    story.append(Spacer(1, 4))
    story.append(Paragraph("<b>1. PrimaryBackupLoanServer.java (extends the Exp 4 BankNode-style server with replication + failover)</b>", ParagraphStyle('SubHead', parent=body_style, fontName='Helvetica-Bold', textColor=colors.HexColor("#006600"))))
    story.append(Spacer(1, 4))

    def create_code_paragraphs(code_text):
        lines = code_text.splitlines()
        flowables = []
        chunk = []
        for line in lines:
            chunk.append(line)
            if len(chunk) >= 40:
                chunk_str = "\n".join(chunk).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>").replace(" ", "&nbsp;")
                flowables.append(Paragraph(chunk_str, code_style))
                flowables.append(Spacer(1, 4))
                chunk = []
        if chunk:
            chunk_str = "\n".join(chunk).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>").replace(" ", "&nbsp;")
            flowables.append(Paragraph(chunk_str, code_style))
        return flowables

    story.extend(create_code_paragraphs(SERVER_CODE))
    story.append(Spacer(1, 8))

    story.append(Paragraph("<b>2. FailoverLoanClient.java (extends the Exp 1/2 LoanClient with automatic reconnect)</b>", ParagraphStyle('SubHead2', parent=body_style, fontName='Helvetica-Bold', textColor=colors.HexColor("#006600"))))
    story.append(Spacer(1, 4))
    story.extend(create_code_paragraphs(CLIENT_CODE))
    story.append(Spacer(1, 12))

    # Section 6: Output screenshot
    story.append(Paragraph("<b>Output screenshot :</b>", section_heading))
    story.append(Spacer(1, 6))

    shot_items = [
        # Real Java Swing Frontend Screenshots
        ("01_Real_Java_UI_Dashboard.png", "Figure 1: BANKCORE Distributed Banking Dashboard (Java Swing) -- Real-time Node Health (Bank Node 01 Primary, Bank Node 02 Backup) and System Metrics"),
        ("02_Real_Java_UI_Loans_Management.png", "Figure 2: BANKCORE Loan Applications Interface (Java Swing) -- Active Loan Registry displaying HL1001 for Gitesh Patil (Rs 62,00,000)"),
        ("03_Real_Java_UI_Loan_Details.png", "Figure 3: BANKCORE Detailed Loan View (Java Swing) -- Verified Application Specification for HL1001 (Customer CUS201 Gitesh Patil)"),
        ("04_Real_Java_UI_Failover_Active.png", "Figure 4: BANKCORE Replication & Failover Monitor (Java Swing) -- Primary Node 01 Crash Simulation, Backup Node 02 Promotion, and Live Sync Stream"),
        
        # Real Terminal Execution Screenshots
        ("01_Primary_Backup_Startup.png", "Figure 5: Primary and Backup Nodes Starting Up (Heartbeat Monitoring Begins on Ports 5000 & 6000)"),
        ("02_Client_Submit_Replicated.png", "Figure 6: Client SUBMIT Request Processed by Primary and Replicated to Backup (Gitesh Patil, Rs 62,00,000)"),
        ("03_Primary_Crash_Failover_Promotion.png", "Figure 7: Primary Crash Simulated -- Backup Detects Timeout (>6000ms) and Self-Promotes to Primary"),
        ("04_Client_Failover_Reconnect.png", "Figure 8: Client Automatically Fails Over to the Newly Promoted Primary (Loan Data HL1001 Intact)")
    ]

    for fname, caption in shot_items:
        img_p = os.path.join(SCREENSHOTS_DIR, fname)
        if os.path.exists(img_p):
            story.append(Paragraph(f"<b>{caption}</b>", body_style))
            story.append(Spacer(1, 4))
            story.append(Image(img_p, width=480, height=270))
            story.append(Spacer(1, 10))

    # Section 7: Conclusion
    story.append(Paragraph("<b>Conclusion:</b> (It should include what the student has learned/ understood from the experiment).", section_heading))
    story.append(Paragraph(
        "In this experiment, we successfully implemented Fault Tolerance in our distributed banking home loan management system using Primary-Backup (Passive) Replication, extending the multithreaded, clock-synchronized, election-capable server built in Experiments 1-4. We demonstrated how the Primary node replicates every state-changing loan request to a Backup node over a dedicated replication channel, how a heartbeat mechanism allows the Backup to detect Primary failure, and how the Backup seamlessly promotes itself to the new Primary without losing any previously acknowledged loan data. We also extended the LoanClient with automatic failover so that customers experience uninterrupted service across a server crash. This experiment highlighted how replication and failure detection together provide high availability and data durability in a distributed banking system.",
        body_style
    ))
    story.append(Spacer(1, 10))

    story.append(Paragraph("<i>Attach research paper based on experiment performed into the classroom (Not for all the practicals)</i>", body_style))
    story.append(Spacer(1, 4))
    story.append(Paragraph("<i>(Students are required to prepare one doc file for each experiment and then upload it in pdf form.)</i>", body_style))

    doc.build(story, onFirstPage=draw_page_border, onLaterPages=draw_page_border)
    print("Exact SPIT Template PDF build complete for Exp 5:", PDF_PATH)

if __name__ == "__main__":
    build_exact_spit_docx()
    build_exact_spit_pdf()
